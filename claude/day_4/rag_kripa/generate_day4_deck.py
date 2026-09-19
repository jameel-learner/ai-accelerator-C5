from pptx import Presentation

prs = Presentation()
prs.core_properties.title = "Day 4: Same RAG, Different Implementation"
prs.core_properties.subject = "Instructor teaching deck"
prs.core_properties.author = "GitHub Copilot"


def add_slide(title, bullets, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    body = slide.shapes.placeholders[1]
    tf = body.text_frame
    tf.clear()
    for idx, item in enumerate(bullets):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.level = 0

    notes_placeholder = slide.notes_slide.placeholders[1]
    notes_placeholder.text_frame.text = notes


add_slide(
    "Day 4: Same RAG, Different Implementation",
    [
        "Day 3 used a framework-based RAG pipeline",
        "Day 4 builds the same logic manually in Python",
        "The concept is the same; the tooling is different",
    ],
    "Instructor notes:\n\nStart by setting the core message. Explain that the architecture did not change—only the abstraction layer changed. Students should understand that a framework like LlamaIndex simplifies the implementation, but the underlying RAG pattern remains the same."
)

add_slide(
    "What was happening in Day 3?",
    [
        "Load documents into the pipeline",
        "Split text into chunks",
        "Embed text and store vectors",
        "Query the vector store and pass context to the LLM",
    ],
    "Instructor notes:\n\nConnect the notebook flow to the conceptual RAG pipeline. Point out that the notebook wrapped this in a higher-level library and made it easy to write in a few commands."
)

add_slide(
    "Where is the same logic in the project?",
    [
        "PDF extraction: src/pdf_parser.py",
        "Chunking: src/chunker.py",
        "Vector storage: src/vector_store.py",
        "Retrieval and generation: src/rag.py",
    ],
    "Instructor notes:\n\nWalk through the file map. Emphasize that the notebook's framework steps are now visible as explicit code in the project. This is the transition from abstraction to implementation."
)

add_slide(
    "Notebook-to-project mapping",
    [
        "LlamaIndex ingestion pipeline -> src/main.py + src/chunker.py",
        "LanceDB vector store -> src/vector_store.py",
        "Query engine -> src/rag.py",
        "LLM provider layer -> src/llm_config.py",
    ],
    "Instructor notes:\n\nThis slide is the key teaching point. Show students that the same conceptual components exist, but are implemented with different tools and code patterns."
)

add_slide(
    "End-to-end flow in the app",
    [
        "Extract text from PDFs",
        "Chunk the text into manageable passages",
        "Embed and store in Chroma",
        "Retrieve relevant docs for the question",
        "Send context + question to the selected LLM",
    ],
    "Instructor notes:\n\nExplain the real-world pipeline as a sequence. This is the same logic the notebook showed, just now implemented in a product-ready structure with Streamlit and a custom orchestration layer."
)

add_slide(
    "Final takeaway",
    [
        "Day 3 = framework-based proof of concept",
        "Day 4 = same RAG architecture as a real app",
        "Different tools, same logic",
        "This is how RAG looks in production-ish Python code",
    ],
    "Instructor notes:\n\nEnd with the message students should remember. The notebook is the conceptual blueprint; the project is the implemented version. They are not separate ideas—they are the same architecture in different forms."
)

out_path = "Day4_RAG_Comparison.pptx"
prs.save(out_path)
print(f"Saved presentation: {out_path}")
