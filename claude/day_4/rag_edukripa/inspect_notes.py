from pptx import Presentation
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
notes = slide.notes_slide
print('type notes slide:', type(notes))
print('dir sample:')
for name in dir(notes):
    if 'text' in name.lower() or 'note' in name.lower() or 'place' in name.lower() or 'slide' in name.lower():
        print(name)
print('has placeholders', hasattr(notes, 'placeholders'))
if hasattr(notes, 'placeholders'):
    print('placeholder count', len(notes.placeholders))
    for i, ph in enumerate(notes.placeholders):
        print('placeholder', i, type(ph), getattr(ph, 'name', None), getattr(ph, 'placeholder_format', None))
        print('has text_frame', hasattr(ph, 'text_frame'))
        if hasattr(ph, 'text_frame'):
            print('text_frame type', type(ph.text_frame))
            print('text_frame contents before', ph.text_frame.text)
        print('attrs', [a for a in dir(ph) if 'text' in a.lower() or 'place' in a.lower()][:20])
